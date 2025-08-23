import os
import csv
import json

from torch import Tensor
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
import openpyxl.chart as xlchart
import openpyxl.utils as xlutil

RESULT_PATH = "/work/u6140562/FoodS2C/FoodImageCode/Results/logs/8_21_aisingle_200/results.json"
OUT_PATH = "/work/u6140562/FoodS2C/FoodImageCode/Results/logs/8_21_aisingle_200"

# Get class names
with open("/work/u6140562/FoodS2C/Database/class.txt", "r") as file :
    cls_names = [line.split()[1] for line in file.readlines()]

# Read result json
with open(RESULT_PATH, "r") as file :
    results = json.load(file)["result"]

def write_csv() :

    # Write metrics
    with open(os.path.join(OUT_PATH, "metrics.csv"), "w", newline="") as file :
        writer = csv.writer(file)
        
        # Write header
        header = ["epoch"] + [key for key, value in results[0].items() \
            if ("train" in key or "valid" in key) and not isinstance(value, (dict, list, Tensor)) and value is not None]
        writer.writerow(header)
        
        # Write metrics
        for res in results :
            writer.writerow(res[key] for key in header)

    list_data_name = [key for key, value in results[0].items() \
        if "valid" in key and isinstance(value, list)]

    # Write metrics of each class
    for name in list_data_name :
        with open(os.path.join(OUT_PATH, f"{name}.csv"), "w", newline="") as file :
            writer = csv.writer(file)
            
            writer.writerow(["epoch"] + cls_names)
            for res in results :
                writer.writerow([res["epoch"]] + res[name])

def write_workbook() :
    
    wb = Workbook()
    scalar_key = [key for key, value in results[0].items() \
        if isinstance(value, (int, float)) and "epoch" not in key.lower()]
    list_key = [key for key, value in results[0].items() \
        if isinstance(value, (list))]
    
    ### Write scalar data ###
    
    ws = wb.active
    
    # Write header
    row, col = 1, 1
    ws.cell(row, col).value = "Epochs"
    for i, key in enumerate(scalar_key, start=col+1) :
        ws.cell(row, i).value = key
    
    # Write step and data
    row += 1
    for i, datas in enumerate(results, start=row) :
        step = datas.get("epoch", i-row)
        ws.cell(i, col).value = step

        # Write data
        for j, k in enumerate(scalar_key, start=col+1) :
            ws.cell(i, j).value = datas[k]

    ### Write list data ###
    
    for key in list_key :
        ws = wb.create_sheet(key)
        
        # Write header
        row, col = 1, 1
        ws.cell(row, col).value = "Epochs"
        for i, name in enumerate(cls_names, start=col+1) :
            ws.cell(row, i).value = name
        
        # Write step and data
        row += 1
        for i, datas in enumerate(results, start=row) :
            step = datas.get("epoch", i-row)
            ws.cell(i, col).value = step

            # Write data
            for j, d in enumerate(datas[key], start=col+1) :
                ws.cell(i, j).value = d
    
    wb.save(os.path.join(OUT_PATH, "out.xlsx"))
    wb.close()

write_workbook()
